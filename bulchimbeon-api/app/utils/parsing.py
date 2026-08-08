"""문서 파싱 (`06 §1` ①) — MD/TXT 그대로, PDF `pypdf`, DOCX `python-docx`.

⚠️ **전부 동기 함수다.** 호출자는 반드시 `run_in_executor` 로 감싼다 (`CLAUDE.md` 룰 8) —
20MB PDF 파싱이 이벤트 루프를 막으면 그동안 모든 요청이 멈춘다.

DOCX 는 헤딩 스타일을 마크다운 헤딩(`##`)으로 변환해 내보낸다. 청킹이 헤딩 경로를
마크다운 문법 하나로만 인식하면 되도록 **입력 포맷의 차이를 여기서 흡수**한다 (`06 §1` ②).
"""

from dataclasses import dataclass
from pathlib import Path

# `05 §4` — 업로드 허용 포맷. 화이트리스트 밖은 400 `UNSUPPORTED_FILE_TYPE` 다.
EXTENSION_TO_MIME: dict[str, str] = {
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

MIME_MARKDOWN = EXTENSION_TO_MIME[".md"]
MIME_TEXT = EXTENSION_TO_MIME[".txt"]
MIME_PDF = EXTENSION_TO_MIME[".pdf"]
MIME_DOCX = EXTENSION_TO_MIME[".docx"]

# 영어 원문 문서를 전제하지만(`04 §2` chunks.content), 업로드가 UTF-8 이 아닐 수 있다.
# 조용히 깨진 글자를 넣지 않도록 후보를 순서대로 시도하고 전부 실패하면 파싱 실패로 본다.
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252")


class DocumentParseError(RuntimeError):
    """파싱 실패. `document_versions.ingest_error` 에 그대로 기록된다 (`06 §1` ①)."""


@dataclass(frozen=True)
class ParsedPage:
    """파싱 산출물 한 덩어리. PDF 만 `page_no` 를 갖는다 (`05 §6` citations 표)."""

    text: str
    page_no: int | None = None


def parse_file(path: Path, mime: str) -> list[ParsedPage]:
    """저장된 파일을 텍스트로 만든다. **동기 함수 — `run_in_executor` 로 부를 것.**"""
    if mime in (MIME_MARKDOWN, MIME_TEXT):
        return [ParsedPage(text=_decode(path.read_bytes()))]
    if mime == MIME_PDF:
        return _parse_pdf(path)
    if mime == MIME_DOCX:
        return [ParsedPage(text=_parse_docx(path))]
    raise DocumentParseError(f"지원하지 않는 mime: {mime}")


def _decode(raw: bytes) -> str:
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError(
        f"텍스트 인코딩을 인식하지 못했습니다 (시도: {', '.join(_TEXT_ENCODINGS)})."
    )


def _parse_pdf(path: Path) -> list[ParsedPage]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        # page_no 는 1부터다 (`05 §6` citations 표).
        return [
            ParsedPage(text=page.extract_text() or "", page_no=index)
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"PDF 파싱 실패: {exc}") from exc


def _parse_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentParseError(f"DOCX 파싱 실패: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = _heading_level(paragraph)
        lines.append(f"{'#' * level} {text}" if level else text)

    # 표는 문단에 들어 있지 않다. 근거 문서에 표가 흔하므로 셀 텍스트도 함께 거둔다.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n\n".join(lines)


def _heading_level(paragraph: object) -> int:
    """`Heading 1`~`Heading 6` → 1~6, 그 외 0.

    스타일 이름은 로캘에 따라 다를 수 있어(`제목 1`) 실패해도 예외를 내지 않는다 —
    헤딩을 못 알아보면 본문으로 취급될 뿐이고, 파싱 자체를 실패시킬 이유는 없다.
    """
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", "") or ""
    prefix = "heading "
    if not name.lower().startswith(prefix):
        return 0
    try:
        level = int(name[len(prefix) :].strip())
    except ValueError:
        return 0
    return level if 1 <= level <= 6 else 0
