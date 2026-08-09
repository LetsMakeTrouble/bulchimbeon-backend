"""M2 문서 테스트용 픽스처 파일 생성 (`05 §4` 의 4개 포맷).

DOCX 는 python-docx 가 쓰기를 지원한다. **PDF 는 쓰기 라이브러리가 없어**(pypdf 는 텍스트를
그려 넣지 못하고 reportlab 은 의존성에 없다) 최소 PDF 를 직접 조립한다 — xref 오프셋까지
맞춰야 pypdf 가 복구 모드 없이 읽는다.
"""

import io
from typing import Any

from httpx import AsyncClient

from tests.helpers import API, Actor

MARKDOWN_SAMPLE = """# Refund Policy

Refunds are handled by the partner success team.

## Standard Refund Window

Refunds are accepted within 30 days of purchase.

## Partial Refunds

Partial refunds follow the same 30-day window.
"""

TEXT_SAMPLE = """# Shipping Policy

Standard shipping takes five business days.

## International

International shipping takes fifteen business days.
"""

PDF_SENTENCE = "Refunds are accepted within 30 days of purchase."

# `fake_llm_provider.embed_failure` 트리거 — **임베딩 입력에 실제로 들어가는** 문자열이어야 한다.
#
# ⚠️ 헤딩 줄(`# Refund Policy`)을 쓰면 안 된다. 인제스트는 헤딩을 뺀 본문을 임베딩하므로
#    (`pipeline/ingest.py` — M-1 캘리브레이션과 조건을 맞추기 위한 것) 헤딩 문자열로는
#    트리거가 걸리지 않는다. 그러면 **실패를 기대한 테스트가 조용히 초록으로 통과한다.**
#    `MARKDOWN_SAMPLE` 의 H1 섹션 본문에서 고른다.
EMBED_FAILURE_TRIGGER = "partner success team"


def build_pdf(sentence: str = PDF_SENTENCE) -> bytes:
    """텍스트 한 줄이 들어 있는 최소 PDF 를 만든다 (pypdf 로 추출 가능)."""
    escaped = sentence.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")

    xref_offset = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode())
    out.write(f"startxref\n{xref_offset}\n%%EOF\n".encode())
    return out.getvalue()


def build_docx(
    title: str = "Support Policy", section: str = "Escalation", body: str = "Escalate within 24h."
) -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading(title, level=1)
    document.add_paragraph("This document describes partner support rules.")
    document.add_heading(section, level=2)
    document.add_paragraph(body)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


SAMPLES: dict[str, tuple[str, bytes, str]] = {
    # 확장자: (파일명, 내용, mime)
    ".md": ("refund-policy.md", MARKDOWN_SAMPLE.encode("utf-8"), "text/markdown"),
    ".txt": ("shipping-policy.txt", TEXT_SAMPLE.encode("utf-8"), "text/plain"),
    ".pdf": ("refund-policy.pdf", build_pdf(), "application/pdf"),
    ".docx": (
        "support-policy.docx",
        build_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
}


async def upload_document(
    client: AsyncClient,
    actor: Actor,
    project_id: str,
    *,
    extension: str = ".md",
    filename: str | None = None,
    content: bytes | None = None,
    title: str | None = None,
    auto_activate: bool | None = None,
    expect_status: int = 201,
) -> dict[str, Any]:
    default_name, default_content, mime = SAMPLES.get(
        extension, ("sample.md", MARKDOWN_SAMPLE.encode("utf-8"), "text/markdown")
    )
    data: dict[str, str] = {}
    if title is not None:
        data["title"] = title
    if auto_activate is not None:
        data["auto_activate"] = "true" if auto_activate else "false"

    response = await client.post(
        f"{API}/projects/{project_id}/documents",
        files={"file": (filename or default_name, content or default_content, mime)},
        data=data,
        headers=actor.headers,
    )
    assert response.status_code == expect_status, response.text
    return response.json()


async def upload_version(
    client: AsyncClient,
    actor: Actor,
    document_id: str,
    *,
    extension: str = ".md",
    content: bytes | None = None,
    auto_activate: bool | None = None,
    expect_status: int = 201,
) -> dict[str, Any]:
    default_name, default_content, mime = SAMPLES[extension]
    data: dict[str, str] = {}
    if auto_activate is not None:
        data["auto_activate"] = "true" if auto_activate else "false"

    response = await client.post(
        f"{API}/documents/{document_id}/versions",
        files={"file": (default_name, content or default_content, mime)},
        data=data,
        headers=actor.headers,
    )
    assert response.status_code == expect_status, response.text
    return response.json()


def version_by_no(document: dict[str, Any], version_no: int) -> dict[str, Any]:
    for version in document["versions"]:
        if version["version_no"] == version_no:
            return version
    raise AssertionError(f"version_no={version_no} 가 없다: {document['versions']}")


def latest_version(document: dict[str, Any]) -> dict[str, Any]:
    """방금 올린 버전.

    ⚠️ 업로드 201 응답의 `active_version` 은 **아직 null** 이다 — `auto_activate=true` 는
    "인제스트가 끝나면 활성화"를 뜻하기 때문이다 (`02 §5` 구현 노트).
    그래서 방금 올린 버전을 가리킬 때는 `active_version` 이 아니라 이 헬퍼를 쓴다.
    """
    return max(document["versions"], key=lambda version: version["version_no"])


async def fetch_document(client: AsyncClient, actor: Actor, document_id: str) -> dict[str, Any]:
    """인제스트가 끝난 뒤의 상태를 다시 읽는다."""
    response = await client.get(f"{API}/documents/{document_id}", headers=actor.headers)
    assert response.status_code == 200, response.text
    return response.json()
